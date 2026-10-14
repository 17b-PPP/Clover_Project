import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_inline_runs(paragraph, text, default_font="TH Sarabun New", default_size=14, is_bold=False, is_italic=False, color=None):
    # Process inline formatting: **bold**, `code`, *italic*
    # Tokenize text
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = default_font
        run._r.get_or_add_rPr().set(qn('w:rFonts'), default_font)
        run.font.size = Pt(default_size)
        if color:
            run.font.color.rgb = color
            
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run.bold = True
            run.text = token[2:-2]
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run.font.name = "Consolas"
            run.font.size = Pt(default_size - 2)
            run.font.color.rgb = RGBColor(199, 37, 78) # Dark pink/red for code
            run.text = token[1:-1]
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run.italic = True
            run.text = token[1:-1]
        else:
            run.bold = is_bold
            run.italic = is_italic
            run.text = token

def build_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    
    # Page Margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'TH Sarabun New'
    normal_style.font.size = Pt(14)
    normal_style.font.color.rgb = RGBColor(33, 37, 41)

    in_code_block = False
    code_block_lines = []
    
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not table_rows:
            in_table = False
            return
        # Parse table rows
        # Filter separator row (e.g. |---|---|)
        parsed_rows = []
        for r in table_rows:
            cells = [c.strip() for c in r.strip('|').split('|')]
            if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                continue
            parsed_rows.append(cells)
            
        if parsed_rows:
            num_cols = max(len(r) for r in parsed_rows)
            table = doc.add_table(rows=len(parsed_rows), cols=num_cols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            
            for row_idx, row_data in enumerate(parsed_rows):
                is_header = (row_idx == 0)
                for col_idx in range(num_cols):
                    cell = table.cell(row_idx, col_idx)
                    cell_text = row_data[col_idx] if col_idx < len(row_data) else ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    
                    if is_header:
                        set_cell_background(cell, "2B547E") # Navy blue header
                        add_inline_runs(p, cell_text, default_size=13, is_bold=True, color=RGBColor(255, 255, 255))
                    else:
                        bg_color = "F8F9FA" if row_idx % 2 == 1 else "FFFFFF"
                        set_cell_background(cell, bg_color)
                        add_inline_runs(p, cell_text, default_size=13, is_bold=False)
                    set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            
            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_after = Pt(6)
            
        table_rows = []
        in_table = False

    def flush_code_block():
        nonlocal in_code_block, code_block_lines
        if not code_block_lines:
            in_code_block = False
            return
        
        # Put code block inside a single-cell table for shaded background & border
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F4F6F8")
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        
        code_text = "\n".join(code_block_lines)
        run = p.add_run(code_text)
        run.font.name = "Consolas"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(36, 41, 46)
        
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(4)
        
        code_block_lines = []
        in_code_block = False

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')
        
        # Check code block
        if line.strip().startswith('```'):
            if in_code_block:
                flush_code_block()
            else:
                if in_table:
                    flush_table()
                in_code_block = True
            i += 1
            continue
            
        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue
            
        # Check table
        if line.strip().startswith('|') and line.strip().endswith('|'):
            in_table = True
            table_rows.append(line)
            i += 1
            continue
        elif in_table:
            flush_table()
            
        # Horizontal rule
        if re.match(r'^-{3,}$', line.strip()):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            # Add a decorative line or divider
            run = p.add_run("―" * 50)
            run.font.color.rgb = RGBColor(180, 180, 180)
            i += 1
            continue
            
        # Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            add_inline_runs(p, line[2:].strip(), default_size=20, is_bold=True, color=RGBColor(24, 76, 120))
            i += 1
            continue
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.keep_with_next = True
            add_inline_runs(p, line[3:].strip(), default_size=17, is_bold=True, color=RGBColor(31, 106, 165))
            i += 1
            continue
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            add_inline_runs(p, line[4:].strip(), default_size=15, is_bold=True, color=RGBColor(44, 62, 80))
            i += 1
            continue
        elif line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            add_inline_runs(p, line[5:].strip(), default_size=14, is_bold=True, color=RGBColor(70, 80, 95))
            i += 1
            continue
            
        # Blockquotes
        if line.startswith('> '):
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            cell.width = Inches(6.5)
            set_cell_background(cell, "EFF6FF") # Soft blue
            set_cell_margins(cell, top=80, bottom=80, left=150, right=150)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            add_inline_runs(p, line[2:].strip(), default_size=13, is_italic=True, color=RGBColor(30, 64, 175))
            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_after = Pt(3)
            i += 1
            continue
            
        # Unordered List
        if re.match(r'^\s*[-*]\s+', line):
            indent_level = (len(line) - len(line.lstrip())) // 2
            text_content = re.sub(r'^\s*[-*]\s+', '', line)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent = Inches(0.25 * (indent_level + 1))
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_inline_runs(p, text_content, default_size=14)
            i += 1
            continue
            
        # Ordered List
        m_num = re.match(r'^\s*(\d+)\.\s+(.*)', line)
        if m_num:
            text_content = m_num.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            add_inline_runs(p, text_content, default_size=14)
            i += 1
            continue
            
        # Blank line
        if not line.strip():
            i += 1
            continue
            
        # Normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        add_inline_runs(p, line.strip(), default_size=14)
        i += 1

    if in_code_block:
        flush_code_block()
    if in_table:
        flush_table()

    doc.save(docx_path)
    print(f"Saved {docx_path} successfully!")

if __name__ == '__main__':
    md_file = r"c:\Users\ASUS\OneDrive\Desktop\Clover Project\PROJECT_DEFENSE_MANUAL.md"
    docx_file1 = r"c:\Users\ASUS\OneDrive\Desktop\Clover Project\คู่มืออธิบายโค้ดและเตรียมสอบโปรเจกต์.docx"
    docx_file2 = r"c:\Users\ASUS\OneDrive\Desktop\Clover Project\PROJECT_DEFENSE_MANUAL.docx"
    
    build_docx(md_file, docx_file1)
    build_docx(md_file, docx_file2)

